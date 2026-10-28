"""Exporta SEGUIMIENTO GRUPOS FOCALES.xlsx -> data/excel-seed.json"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(r"c:\Users\CAMPUSLANDS\Downloads\SEGUIMIENTO GRUPOS FOCALES.xlsx")
OUT = ROOT / "data" / "excel-seed.json"


def norm(s):
    if s is None:
        return None
    s = re.sub(r"\s+", " ", str(s).strip())
    return s or None


def split_resp(s):
    if not s:
        return []
    return [p for p in (norm(x) for x in re.split(r"[\n/]+", str(s))) if p]


def term_pts(t):
    t = (t or "").strip().lower()
    if "cr" in t:
        return 20, "CRITICO"
    if "pos" in t:
        return 100, "POSITIVO"
    return 60, "MEDIO"


def estado_norm(e):
    e = (e or "").strip().lower()
    if "cerr" in e:
        return "CERRADO"
    if "proceso" in e:
        return "EN_PROCESO"
    return "ABIERTO"


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    hallazgos = []
    areas, cats, subs, resps = set(), set(), set(), set()

    for sheet_name, ciclo_nombre in [("CICLO 1", "Ciclo 1"), ("CICLO 2", "Ciclo 2")]:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            fecha = ws.cell(r, 2).value
            if fecha is None:
                continue
            area = norm(ws.cell(r, 3).value)
            cat = norm(ws.cell(r, 4).value)
            sub = norm(ws.cell(r, 5).value)
            term_raw = norm(ws.cell(r, 6).value)
            obs = ws.cell(r, 7).value
            comp = ws.cell(r, 8).value
            resp_raw = ws.cell(r, 9).value
            estado = estado_norm(ws.cell(r, 10).value)
            avances = ws.cell(r, 11).value
            calif = ws.cell(r, 12).value
            pts, term = term_pts(term_raw)
            if calif is not None:
                try:
                    pts = int(float(calif))
                except Exception:
                    pass
            if isinstance(fecha, datetime):
                fecha_s = fecha.date().isoformat()
            elif isinstance(fecha, date):
                fecha_s = fecha.isoformat()
            else:
                fecha_s = str(fecha)[:10]
            responsables = split_resp(resp_raw)
            areas.add(area)
            cats.add(cat)
            if sub:
                subs.add((cat, sub))
            for rr in responsables:
                resps.add(rr)
            hallazgos.append(
                {
                    "ciclo": ciclo_nombre,
                    "fecha": fecha_s,
                    "area": area,
                    "categoria": cat,
                    "subcategoria": sub,
                    "termometro": term,
                    "observacion": str(obs).strip() if obs else "",
                    "compromiso": str(comp).strip() if comp else "",
                    "responsables": responsables,
                    "estado": estado,
                    "avances": str(avances).strip() if avances else "",
                    "calificacion": pts,
                }
            )

    out = {
        "areas": sorted(a for a in areas if a),
        "categorias": sorted(c for c in cats if c),
        "subcategorias": [{"categoria": c, "nombre": s} for c, s in sorted(subs)],
        "responsables": sorted(resps),
        "escala": {
            "termometro": {"POSITIVO": 100, "MEDIO": 60, "CRITICO": 20},
            "estados_area": [
                {"min": 85, "max": 100, "estado": "Excelente"},
                {"min": 70, "max": 84, "estado": "Aceptable / Estable"},
                {"min": 50, "max": 69, "estado": "En Riesgo / Alerta"},
                {"min": 0, "max": 49, "estado": "Crítico / Vulnerable"},
            ],
        },
        "hallazgos": hallazgos,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK {len(hallazgos)} hallazgos -> {OUT}")


if __name__ == "__main__":
    main()
